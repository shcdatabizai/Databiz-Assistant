#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
신한카드 제안서 공통 디자인 시스템 (python-pptx)

좌표계
    슬라이드를 680 x 383 "단위"로 잡고 I() 로 인치(EMU)로 환산한다.
    16:9 기준 13.3 x 7.5 in 에 대응하며, 모든 페이지 함수는 이 단위로 작성한다.
    y 좌표는 '텍스트 기준선(baseline)' 기준이다.

레이아웃 관례
      0 ~  44   헤더 바 (번호 / 분모 / 제목 / 로고)
     56 ~  98   헤드 메시지 (14px 굵게 + 11px 보조)
    116          섹션 라벨 ①②③ (sec)
    116 ~ 350   본문 영역
    360 ~ 378   주석 (note)
    좌우 여백 28, 2단 구성 시 좌 28~330 / 우 352~652

사용
    from shinhan_deck import *
    prs = new_deck()
    sl = blank(prs)
    frame(sl, "01", "07", "데이터 개요")
    head(sl, "메인 메시지", "보조 설명")
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
from lxml import etree

# ─────────────────────────────────────── 좌표 / 서체
UNIT = 13.3 / 680
I = lambda v: Emu(int(v * UNIT * 914400))
P = lambda px: Pt(round(px * 1.408, 1))      # 단위 → 폰트 포인트
W = lambda px: Pt(round(px * 1.408, 2))      # 단위 → 선 두께

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
# 원신한은 굵기별로 패밀리가 분리돼 있어 bold 속성 대신 패밀리를 바꾼다.
# 영문 패밀리명을 쓰는 이유: 한글명은 한국어 Windows 밖에서 인식되지 않을 수 있다.
FONT_BOLD, FONT_MED = "OneShinhan Bold", "OneShinhan Medium"

# ─────────────────────────────────────── 색상 토큰
C = dict(
    blue="0046FF",      # 신한 블루 (강조 / 헤더)
    navy="0B2E6F",      # 제목 텍스트
    mid="3D4A66",       # 본문 텍스트
    gray="7A88A6",      # 보조 / 주석
    lgray="B4BDCE",     # 축 라벨
    pale="C3CAD8",      # 빈 값(–)
    b1="3D74FF", b2="7FA3FF", b3="B3CCFF", b4="9DBBFF",   # 블루 스케일
    cpat="6A94FF",      # 별도 표기용(쿠팡 등)
    tint="E6EEFF",      # 강조 배경
    soft="F2F6FF",      # 카드 배경
    card="C9D6F0",      # 슬라이드 테두리
    line="E1E8F5",      # 표 구분선
    sub="5A6478",       # 헤드 보조문
    hsub="8FB2FF",      # 헤더 분모 / 표지 보조
    white="FFFFFF",
    red="B03030",       # 경고성 주석
    redl="C25E5E",      # ILLUSTRATIVE 도장 글자
    redb="E09090",      # 도장 테두리
    pinkbg="FDF6F6",    # 예시 블록 배경
    pinkbd="E8C4C4",    # 예시 블록 테두리
    pinkln="E8D0D0",    # 예시 블록 내부선
)

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO_KO = os.path.join(HERE, "assets", "logo_ko_white.png")
LOGO_EN = os.path.join(HERE, "assets", "logo_en_white.png")


# ─────────────────────────────────────── 문서 / 슬라이드
def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 13.333 x 7.5 in
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ─────────────────────────────────────── 기본 드로잉
def _set_face(run, face):
    """latin / ea(한글) / cs 를 모두 같은 패밀리로 지정해야 혼용이 안 생긴다."""
    run.font.name = face
    rpr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rpr.find("{%s}%s" % (A_NS, tag))
        if el is None:
            el = etree.SubElement(rpr, "{%s}%s" % (A_NS, tag))
        el.set("typeface", face)


def txt(sl, x, y, text, size, color, bold=False, align="l", w=260, rotate=None, spacing=None):
    """y 는 baseline. align: l(기본) / r(x가 오른쪽 끝) / c(x가 중앙)"""
    left = x - w if align == "r" else (x - w / 2 if align == "c" else x)
    tb = sl.shapes.add_textbox(I(left), I(y - size * 1.20), I(w), I(size * 1.65))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[align]
    r = p.add_run()
    r.text = text
    r.font.size = P(size)
    r.font.color.rgb = RGBColor.from_string(color)
    _set_face(r, FONT_BOLD if bold else FONT_MED)
    if spacing:
        r._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    if rotate:
        tb.rotation = rotate
    return tb


def rect(sl, x, y, w, h, fill=None, rx=None, line=None, dash=False):
    shp = sl.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rx else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if rx:
        shp.adjustments[0] = min(0.5, rx / min(w, h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line[0])
        shp.line.width = W(line[1])
        if dash:
            shp.line.dash_style = 4
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def hline(sl, x1, x2, y, color, width=0.6):
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y), I(x2), I(y))
    cn.line.color.rgb = RGBColor.from_string(color)
    cn.line.width = W(width)
    return cn


def seg(sl, x1, y1, x2, y2, color, width=1.3):
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    cn.line.color.rgb = RGBColor.from_string(color)
    cn.line.width = W(width)
    return cn


def polyline(sl, xs, ys, color, width=1.3):
    for i in range(len(xs) - 1):
        seg(sl, xs[i], ys[i], xs[i + 1], ys[i + 1], color, width)


def bar(sl, x, y, w, h, color, rx=None):
    """0 이하 길이는 그리지 않는다 (값이 0인 막대로 인한 점 노이즈 방지)"""
    if w > 0.5:
        rect(sl, x, y, w, h, fill=color, rx=rx)


def dot(sl, cx, cy, r, fill=None, hollow=False):
    o = sl.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - r), I(cy - r), I(2 * r), I(2 * r))
    if hollow:
        o.fill.background()
        o.line.color.rgb = RGBColor.from_string(C["lgray"])
        o.line.width = W(0.9)
    else:
        o.fill.solid()
        o.fill.fore_color.rgb = RGBColor.from_string(fill)
        o.line.fill.background()
    o.shadow.inherit = False
    return o


# ─────────────────────────────────────── 페이지 골격
def frame(sl, num, total, title):
    """헤더 바 + 페이지 번호 + 로고. total 은 '07' 같은 문자열"""
    rect(sl, 0, 0, 680, 383, fill=C["white"], line=(C["card"], 0.6))
    rect(sl, 0, 0, 680, 44, fill=C["blue"])
    txt(sl, 28, 31, num, 20, C["white"], bold=True, w=44)
    txt(sl, 54, 31, f"／ {total}", 11, C["hsub"], w=60)
    txt(sl, 92, 30, title, 15, C["white"], bold=True, w=460)
    if os.path.exists(LOGO_EN):
        sl.shapes.add_picture(LOGO_EN, I(566), I(14), I(88), I(15))


def head(sl, line1, line2):
    """헤드 메시지. line1 은 한 줄에 들어가야 한다(14px 기준 약 40자)"""
    rect(sl, 28, 56, 4, 42, fill=C["blue"])
    txt(sl, 42, 72, line1, 14, C["navy"], bold=True, w=600)
    txt(sl, 42, 89, line2, 11, C["sub"], w=600)


def sec(sl, x, x_end, y, label):
    """섹션 라벨 + 밑줄"""
    txt(sl, x, y, label, 11, C["blue"], bold=True, w=x_end - x)
    hline(sl, x, x_end, y + 6, C["blue"], 1)


def note(sl, y, text, color=None, size=10):
    txt(sl, 28, y, text, size, color or C["gray"], w=624)


def illu_block(sl, x, y, w, h, title):
    """예시(illustrative) 영역: 분홍 점선 블록 + 붉은 제목"""
    rect(sl, x, y, w, h, fill=C["pinkbg"], rx=4, line=(C["pinkbd"], 0.9), dash=True)
    txt(sl, x + 12, y + 17, title, 11, C["red"], bold=True, w=w - 24)


def stamp(sl, cx, cy, text="ILLUSTRATIVE"):
    """기울어진 ILLUSTRATIVE 도장"""
    w, h = 116, 22
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            I(cx - w / 2), I(cy - h / 2), I(w), I(h))
    s.adjustments[0] = 2 / h
    s.fill.background()
    s.line.color.rgb = RGBColor.from_string(C["redb"])
    s.line.width = W(1.2)
    s.shadow.inherit = False
    s.rotation = 9
    txt(sl, cx, cy + 4, text, 11, C["redl"], align="c", w=140, rotate=9, spacing=2)


# ─────────────────────────────────────── 표지 / 목차
def cover(prs, title1, title2, subtitle, org, date, chip="DATA PROPOSAL"):
    sl = blank(prs)
    rect(sl, 0, 0, 680, 383, fill=C["navy"])
    rect(sl, 64, 62, 150, 26, fill="17408A", rx=13)
    txt(sl, 139, 79, chip, 11, "B3CCFF", align="c", w=150)
    rect(sl, 64, 152, 44, 4, fill=C["b1"], rx=2)
    txt(sl, 64, 198, title1, 24, C["white"], bold=True, w=560)
    txt(sl, 64, 234, title2, 24, C["white"], bold=True, w=560)
    txt(sl, 64, 266, subtitle, 12, C["hsub"], w=460)
    if os.path.exists(LOGO_KO):
        sl.shapes.add_picture(LOGO_KO, I(64), I(318), I(97), I(26))
    txt(sl, 616, 336, f"{date}  |  {org}", 11, C["hsub"], align="r", w=300)
    return sl


def toc(prs, items, heading="Table of Contents"):
    """items: [(번호, 제목, 설명), ...] — 행 간격은 개수에 따라 자동"""
    sl = blank(prs)
    rect(sl, 0, 0, 680, 383, fill=C["white"], line=(C["card"], 0.6))
    rect(sl, 0, 0, 680, 44, fill=C["blue"])
    txt(sl, 28, 30, heading, 15, C["white"], bold=True, w=320)
    if os.path.exists(LOGO_EN):
        sl.shapes.add_picture(LOGO_EN, I(566), I(14), I(88), I(15))
    n = len(items)
    top = 78
    pitch = min(58, (340 - top) / max(n - 1, 1))
    for i, (num, title, desc) in enumerate(items):
        y = top + i * pitch
        rect(sl, 44, y - 12, 34, 26, fill=C["tint"], rx=5)
        txt(sl, 61, y + 5, num, 12, C["blue"], bold=True, align="c", w=34)
        txt(sl, 96, y, title, 13, C["navy"], bold=True, w=300)
        txt(sl, 96, y + 15, desc, 10, C["gray"], w=520)
        hline(sl, 44, 636, y + 25, C["line"], 0.6)
    return sl
